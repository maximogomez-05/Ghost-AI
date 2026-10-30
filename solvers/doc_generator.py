import os
import sys
import json
import base64
import subprocess
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_shading(cell, color_hex):
    shd_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))

def format_cell_text(cell, text, bold=False, italic=False, color="000000", size=10, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color != "000000":
        r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)

def create_document(title, subtitle, exercise_text, sections, output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Title
    t_p = doc.add_paragraph()
    t_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_p.paragraph_format.space_after = Pt(2)
    r_t = t_p.add_run(title)
    r_t.font.size = Pt(16)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Subtitle
    if subtitle:
        s_p = doc.add_paragraph()
        s_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s_p.paragraph_format.space_after = Pt(14)
        r_s = s_p.add_run(subtitle)
        r_s.font.size = Pt(11)
        r_s.italic = True
        r_s.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # Enunciado box
    if exercise_text:
        e_p = doc.add_paragraph()
        e_p.paragraph_format.space_after = Pt(12)
        r_lbl = e_p.add_run("Consigna del Ejercicio:\n")
        r_lbl.bold = True
        r_lbl.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        e_p.add_run(exercise_text)

    # Process structured sections
    for sec in sections:
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(14)
        h_p.paragraph_format.space_after = Pt(4)
        r_h = h_p.add_run(sec.get("heading", ""))
        r_h.font.size = Pt(13)
        r_h.bold = True
        r_h.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        content = sec.get("text", "")
        if content:
            cp = doc.add_paragraph()
            cp.paragraph_format.space_after = Pt(6)
            cp.add_run(content)

        # Image if any
        img_path = sec.get("image_path")
        if img_path and os.path.exists(img_path):
            ip = doc.add_paragraph()
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.paragraph_format.space_before = Pt(6)
            ip.paragraph_format.space_after = Pt(4)
            ip.add_run().add_picture(img_path, width=Inches(5.8))
            if sec.get("image_caption"):
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_p.paragraph_format.space_after = Pt(8)
                r_cap = cap_p.add_run(sec.get("image_caption"))
                r_cap.font.size = Pt(9.5)
                r_cap.italic = True
                r_cap.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

        # Table if any
        table_data = sec.get("table")
        if table_data:
            headers = table_data.get("headers", [])
            rows = table_data.get("rows", [])
            if headers:
                tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                tbl.style = 'Table Grid'
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for c_idx, h_text in enumerate(headers):
                    cell = tbl.cell(0, c_idx)
                    format_cell_text(cell, h_text, bold=True, color="FFFFFF", size=9.5)
                    set_cell_shading(cell, "1E3A8A")
                for r_idx, row in enumerate(rows, start=1):
                    is_opt = table_data.get("highlight_row") == r_idx
                    for c_idx, val in enumerate(row):
                        cell = tbl.cell(r_idx, c_idx)
                        format_cell_text(cell, val, bold=is_opt, size=9.5)
                        if is_opt:
                            set_cell_shading(cell, "DCFCE7")
                        elif r_idx % 2 == 1:
                            set_cell_shading(cell, "F9FAFB")

    doc.save(output_path)
    return output_path

def convert_to_pdf(docx_path):
    outdir = os.path.dirname(docx_path)
    soffice_candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.com",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        "soffice"
    ]
    soffice_bin = None
    for cand in soffice_candidates:
        if os.path.exists(cand) or cand == "soffice":
            soffice_bin = cand
            break
            
    if not soffice_bin:
        return None
        
    try:
        cmd = [soffice_bin, "--headless", "--convert-to", "pdf", docx_path, "--outdir", outdir]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        base = os.path.splitext(docx_path)[0]
        pdf_path = base + ".pdf"
        if os.path.exists(pdf_path):
            return pdf_path
    except Exception as e:
        print(f"PDF conversion error: {e}")
    return None

def upload_to_temp(file_path):
    if not os.path.exists(file_path):
        return None
    try:
        cmd = ["curl.exe", "-F", f"file=@{file_path}", "https://temp.sh/upload"]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=15)
        out = res.stdout.strip()
        lines = out.split("\n")
        for line in lines:
            line = line.strip()
            if line.startswith("https://temp.sh/"):
                return line
    except Exception as e:
        print(f"Upload error: {e}")
    return None

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--test":
        test_out = r"c:\Users\ariel\testeos\test_doc_generator.docx"
        create_document("Documento de Prueba", "Resolución Oficial Ghost", "Consigna de prueba", [], test_out)
        print("Generated:", test_out)
